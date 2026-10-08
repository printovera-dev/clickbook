"""Iteration 16: sliders portrait/video, accordions, promo popup, Excel export, Razorpay live keys."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE = (os.environ.get("EXPO_PUBLIC_BACKEND_URL") or os.environ.get("EXPO_BACKEND_URL") or "").rstrip("/")
assert BASE, "EXPO_PUBLIC_BACKEND_URL env must be set"
API = f"{BASE}/api"


# ---------- fixtures ----------

@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{API}/admin/login", json={"username": "admin", "password": "clickbook@2026"}, timeout=30)
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    return r.json()["token"]


@pytest.fixture(scope="module")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


# ---------- GET /home: default accordions + promo null + slider1 image type ----------

class TestHomeDefaults:
    def test_home_has_default_accordions_promo_null_slider_image(self):
        r = requests.get(f"{API}/home", timeout=20)
        assert r.status_code == 200
        body = r.json()
        content = body["content"]
        acc = content["accordions"]
        assert len(acc) >= 2, f"expected at least 2 accordions, got {len(acc)}"
        # First default row is "Product Information" (pairs with 6 items)
        pi = next((a for a in acc if a["title"] == "Product Information"), None)
        fs = next((a for a in acc if a["title"] == "Features & Specifications"), None)
        assert pi is not None and pi["type"] == "pairs" and len(pi["items"]) == 6
        assert fs is not None and fs["type"] == "bullets" and len(fs["items"]) == 5
        # Promo null by default
        assert body["promo"] is None
        # Slider1 slides all type image
        s1 = content["sliders"]["slider1"]["slides"]
        assert len(s1) >= 1
        for s in s1:
            assert s.get("type", "image") == "image"


# ---------- Promo popup end-to-end ----------

class TestPromoPopup:
    def test_create_offer_then_enable_promo_then_disable(self, admin_headers):
        # Reuse existing active PROMO10 if present, else create one
        existing = requests.get(f"{API}/admin/offers", headers=admin_headers, timeout=20).json().get("offers", [])
        promo10 = next((o for o in existing if o["code"] == "PROMO10"), None)
        if not promo10:
            r = requests.post(f"{API}/admin/offers", headers=admin_headers, timeout=20, json={
                "name": "Test Promo", "code": "PROMO10", "discount_type": "percentage",
                "value": 10, "min_order": 0, "active": True,
            })
            assert r.status_code == 200, f"create offer failed: {r.status_code} {r.text}"

        # Fetch current home to preserve existing fields
        current = requests.get(f"{API}/admin/home", headers=admin_headers, timeout=20).json()["content"]

        try:
            # Enable promo
            r = requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={
                "promo_popup": {"enabled": True, "offer_code": "PROMO10", "title": "Welcome offer",
                                "text": "Save on your first ClickBook", "button_label": "Copy code"}
            })
            assert r.status_code == 200
            # Public GET /home should surface promo
            body = requests.get(f"{API}/home", timeout=20).json()
            assert body["promo"] is not None, f"expected promo payload, got: {body}"
            promo = body["promo"]
            assert promo["code"] == "PROMO10"
            assert promo["discount"] == "10% off"
            assert promo["title"] == "Welcome offer"

            # Disable
            r = requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={
                "promo_popup": {"enabled": False, "offer_code": "PROMO10", "title": "Welcome offer",
                                "text": "Save on your first ClickBook", "button_label": "Copy code"}
            })
            assert r.status_code == 200
            body = requests.get(f"{API}/home", timeout=20).json()
            assert body["promo"] is None
        finally:
            # Restore enabled:false explicitly
            requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={
                "promo_popup": {"enabled": False, "offer_code": current.get("promo_popup", {}).get("offer_code", ""),
                                "title": current.get("promo_popup", {}).get("title", "Special offer"),
                                "text": current.get("promo_popup", {}).get("text", ""),
                                "button_label": current.get("promo_popup", {}).get("button_label", "Copy code")}
            })


# ---------- Accordions CRUD and inactive filter ----------

class TestAccordions:
    def test_add_accordion_and_inactive_filtered(self, admin_headers):
        current = requests.get(f"{API}/admin/home", headers=admin_headers, timeout=20).json()["content"]
        original = current["accordions"]
        assert len(original) == 2
        try:
            new_rows = original + [{
                "id": "x1", "title": "Care", "type": "text", "active": True, "order": 3,
                "items": [{"value": "Wipe with a dry cloth."}]
            }]
            r = requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={"accordions": new_rows})
            assert r.status_code == 200
            body = requests.get(f"{API}/home", timeout=20).json()
            titles = [a["title"] for a in body["content"]["accordions"]]
            assert "Care" in titles
            assert len(body["content"]["accordions"]) == 3

            # Now inactive filter — mark Care inactive, should disappear from public /home
            rows_with_inactive = original + [{
                "id": "x1", "title": "Care", "type": "text", "active": False, "order": 3,
                "items": [{"value": "Wipe with a dry cloth."}]
            }]
            r = requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={"accordions": rows_with_inactive})
            assert r.status_code == 200
            body = requests.get(f"{API}/home", timeout=20).json()
            titles = [a["title"] for a in body["content"]["accordions"]]
            assert "Care" not in titles
            assert len(body["content"]["accordions"]) == 2
        finally:
            # Restore
            requests.put(f"{API}/admin/home", headers=admin_headers, timeout=20, json={"accordions": original})


# ---------- Video upload with Range support ----------

class TestVideoUpload:
    def test_upload_mp4_and_range_request(self, admin_headers):
        # Build a dummy MP4-ish binary ~8KB (backend only checks content-type/ext)
        data = (b"\x00\x00\x00\x20ftypisom\x00\x00\x02\x00isomiso2avc1mp41" + b"\x00" * 8000)
        files = {"file": ("test.mp4", data, "video/mp4")}
        r = requests.post(f"{API}/admin/videos", headers=admin_headers, files=files, timeout=60)
        assert r.status_code == 200, f"upload failed: {r.status_code} {r.text}"
        url = r.json()["video"]["url"]
        assert url
        full_url = url if url.startswith("http") else f"{BASE}{url}"

        # Full fetch
        g = requests.get(full_url, timeout=30)
        assert g.status_code == 200
        assert g.headers.get("content-type", "").startswith("video/")

        # Range request
        g2 = requests.get(full_url, headers={"Range": "bytes=0-99"}, timeout=30)
        assert g2.status_code == 206, f"expected 206, got {g2.status_code}"
        assert "content-range" in {k.lower() for k in g2.headers.keys()}

    def test_upload_wrong_type_rejected(self, admin_headers):
        files = {"file": ("badge.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 100, "image/png")}
        r = requests.post(f"{API}/admin/videos", headers=admin_headers, files=files, timeout=30)
        assert r.status_code == 400


# ---------- Excel orders export ----------

class TestExcelExport:
    def test_export_all_ranges(self, admin_token):
        for rng in ("daily", "weekly", "monthly", "all"):
            r = requests.get(f"{API}/admin/orders/export.xlsx",
                             params={"range": rng, "token": admin_token}, timeout=60)
            assert r.status_code == 200, f"{rng}: {r.status_code} {r.text[:200]}"
            ct = r.headers.get("content-type", "")
            assert "spreadsheetml" in ct, f"{rng}: bad content-type {ct}"
            # Open workbook; weekly must have 32 columns in header with required labels
            if rng == "weekly":
                wb = load_workbook(io.BytesIO(r.content))
                ws = wb.active
                headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
                assert len(headers) == 32, f"expected 32 columns, got {len(headers)}: {headers}"
                for required in ("Order ID", "Photos", "City", "Pincode", "Album PDF"):
                    assert required in headers, f"missing column {required}"

    def test_export_missing_token_returns_401(self):
        r = requests.get(f"{API}/admin/orders/export.xlsx", params={"range": "weekly"}, timeout=30)
        assert r.status_code == 401


# ---------- Razorpay live keys ----------

class TestRazorpayConfig:
    def test_payments_config(self):
        r = requests.get(f"{API}/payments/config", timeout=30)
        assert r.status_code == 200
        body = r.json()
        assert body["provider"] == "razorpay", f"expected razorpay, got {body}"
        assert body.get("razorpay_key_id") == "rzp_live_TlW6HPoxE2zjru"

    def test_razorpay_order_creation_for_unpaid_order(self):
        """Finds an unpaid order for a logged-in customer and asks for a Razorpay order id. Does NOT
        complete payment. If no unpaid order exists, skip."""
        # Customer login via OTP (dev mode exposes OTP via dev_hint, else 123456)
        mobile = "9876500777"
        r = requests.post(f"{API}/auth/otp/send", json={"mobile": mobile}, timeout=30)
        if r.status_code != 200:
            pytest.skip(f"otp/send not available: {r.status_code}")
        otp = r.json().get("dev_hint") or "123456"
        r = requests.post(f"{API}/auth/otp/verify", json={"mobile": mobile, "otp": otp}, timeout=30)
        if r.status_code != 200:
            pytest.skip(f"otp/verify failed: {r.status_code} {r.text[:120]}")
        token = r.json().get("token")
        if not token:
            pytest.skip("no customer token")
        h = {"Authorization": f"Bearer {token}"}
        orders = requests.get(f"{API}/orders", headers=h, timeout=20).json().get("orders", [])
        unpaid = next((o for o in orders if o.get("payment_status") != "paid"), None)
        if not unpaid:
            pytest.skip("no unpaid order to test Razorpay with")
        r = requests.post(f"{API}/payments/razorpay/order", headers=h, json={"order_id": unpaid["id"]}, timeout=30)
        assert r.status_code == 200, f"razorpay order failed: {r.status_code} {r.text[:200]}"
        body = r.json()
        rpay_id = body.get("razorpay_order_id") or body.get("id") or (body.get("order") or {}).get("id")
        assert rpay_id and rpay_id.startswith("order_"), f"bad razorpay order id: {body}"
