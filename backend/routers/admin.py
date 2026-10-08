"""Admin console: login, dashboard, orders, covers, offers, settings, process bots, customers, PDF."""
import io
from datetime import datetime, timezone, timedelta
from pathlib import Path
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Header
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel
from typing import Optional, List, Dict, Any

from core import (db, now_iso, new_id, get_current_admin, get_settings_doc, verify_password, hash_password,
                  make_admin_token, admin_from_token)
from storage_manager import save_original, public_url, list_dir_files, zip_dir, _put
from production import schedule_production_package, production_build_is_stale, STALE_BUILD_MESSAGE
from routers.notifications import notify_order_status

router = APIRouter(prefix="/admin", tags=["admin"])

PRODUCTION_STATUSES = {"processing", "printing", "packaging", "out_for_delivery", "delivered", "cancelled", "refunded"}


class AdminLogin(BaseModel):
    username: str
    password: str


class CoverIn(BaseModel):
    name: str
    description: Optional[str] = ""
    image_url: str
    style: Optional[str] = None  # signature | classic | editorial
    active: bool = True
    display_order: int = 0


class OfferIn(BaseModel):
    name: str
    code: str
    discount_type: str  # percentage | fixed
    value: float
    min_order: float = 0
    max_discount: Optional[float] = None
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    active: bool = True
    public: bool = True
    usage_limit: Optional[int] = None


class SettingsUpdate(BaseModel):
    price_per_sheet: Optional[float] = None
    gst_percent: Optional[float] = None
    min_sheets: Optional[int] = None
    max_sheets: Optional[int] = None
    gift_wrap_fee: Optional[float] = None
    support_whatsapp: Optional[str] = None


class StatusUpdate(BaseModel):
    production_status: str
    note: Optional[str] = ""
    tracking_number: Optional[str] = None


class BotUpdate(BaseModel):
    label: Optional[str] = None
    message: Optional[str] = None
    icon: Optional[str] = None
    image_url: Optional[str] = None
    active: Optional[bool] = None


@router.post("/login")
async def admin_login(payload: AdminLogin):
    admin = await db.admins.find_one({"username": payload.username, "active": {"$ne": False}}, {"_id": 0})
    if not admin or not verify_password(payload.password, admin.get("password_hash")):
        raise HTTPException(401, "Incorrect username or password")
    token, expires_in = make_admin_token(admin["username"])
    await db.admins.update_one({"username": admin["username"]}, {"$set": {"last_login_at": now_iso()}})
    return {"token": token, "expires_in": expires_in, "admin": {"username": admin["username"], "role": admin.get("role", "super")}}


@router.get("/me")
async def admin_me(admin: dict = Depends(get_current_admin)):
    return {"admin": {"username": admin["username"], "role": admin.get("role", "super")}}


class PasswordChange(BaseModel):
    current_password: str
    new_password: str


@router.put("/password")
async def admin_change_password(payload: PasswordChange, admin: dict = Depends(get_current_admin)):
    doc = await db.admins.find_one({"username": admin["username"]}, {"_id": 0, "password_hash": 1})
    if not verify_password(payload.current_password, (doc or {}).get("password_hash")):
        raise HTTPException(400, "Current password is incorrect")
    if len(payload.new_password) < 8:
        raise HTTPException(400, "New password must be at least 8 characters")
    await db.admins.update_one({"username": admin["username"]}, {"$set": {"password_hash": hash_password(payload.new_password)}})
    return {"success": True}


@router.get("/dashboard")
async def admin_dashboard(admin: dict = Depends(get_current_admin)):
    orders = await db.orders.find({}, {"_id": 0}).to_list(1000)
    paid = [o for o in orders if o.get("payment_status") == "paid"]
    revenue = sum(o["price"]["total"] for o in paid)
    by_status: dict = {}
    for o in orders:
        st = o.get("production_status", "processing")
        by_status[st] = by_status.get(st, 0) + 1
    # Revenue / volume by day for the last 14 days (dashboard chart)
    from datetime import datetime, timedelta, timezone
    today = datetime.now(timezone.utc).date()
    days = [(today - timedelta(days=i)).isoformat() for i in range(13, -1, -1)]
    daily = {d: {"date": d, "revenue": 0.0, "orders": 0} for d in days}
    for o in paid:
        d = (o.get("paid_at") or o.get("created_at") or "")[:10]
        if d in daily:
            daily[d]["revenue"] += o["price"]["total"]
            daily[d]["orders"] += 1
    return {
        "total_orders": len(orders),
        "paid_orders": len(paid),
        "pending_payments": len([o for o in orders if o.get("payment_status") != "paid"]),
        "revenue": round(revenue, 2),
        "sheets_sold": sum(o.get("sheets", 0) for o in paid),
        "by_status": by_status,
        "customers": await db.customers.count_documents({}),
        "albums": await db.albums.count_documents({}),
        "drafts": await db.albums.count_documents({"status": "draft"}),
        "daily": [daily[d] for d in days],
        "recent_orders": sorted(orders, key=lambda o: o.get("created_at", ""), reverse=True)[:5],
    }


@router.get("/payments")
async def admin_payments(admin: dict = Depends(get_current_admin)):
    orders = await db.orders.find({}, {"_id": 0, "id": 1, "order_no": 1, "client_name": 1, "customer_snapshot": 1, "sheets": 1,
                                       "price": 1, "payment_status": 1, "payment_method": 1, "payment_id": 1, "paid_at": 1,
                                       "created_at": 1, "coupon_code": 1, "gift_wrap": 1}).sort("created_at", -1).to_list(500)
    paid = [o for o in orders if o.get("payment_status") == "paid"]
    return {"payments": orders,
            "summary": {"collected": round(sum(o["price"]["total"] for o in paid), 2), "paid": len(paid),
                        "pending": len(orders) - len(paid),
                        "pending_amount": round(sum(o["price"]["total"] for o in orders if o.get("payment_status") != "paid"), 2)}}


@router.get("/albums")
async def admin_albums(admin: dict = Depends(get_current_admin), status: Optional[str] = None):
    q = {"status": status} if status else {}
    albums = await db.albums.find(q, {"_id": 0, "pages": 0, "design_versions": 0, "album_snapshot": 0}).sort("updated_at", -1).to_list(500)
    cust_ids = list({a.get("customer_id") for a in albums if a.get("customer_id")})
    customers = {c["id"]: c for c in await db.customers.find({"id": {"$in": cust_ids}}, {"_id": 0, "id": 1, "name": 1, "mobile": 1}).to_list(500)}
    out = []
    for a in albums:
        photos = a.pop("photos", []) or []
        cover_photo = next((p for p in photos if p.get("id") == (a.get("cover_design") or {}).get("photo_id")), photos[0] if photos else None)
        c = customers.get(a.get("customer_id"), {})
        out.append({**a, "photo_count": len(photos), "customer_name": c.get("name"), "customer_mobile": c.get("mobile"),
                    "cover_thumbnail_url": (cover_photo or {}).get("thumbnail_url")})
    return {"albums": out}


# ---- Orders ----

@router.get("/orders")
async def admin_orders(admin: dict = Depends(get_current_admin), status: Optional[str] = None):
    q = {"production_status": status} if status else {}
    orders = await db.orders.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    return {"orders": orders}


@router.put("/orders/{order_id}/status")
async def update_order_status(order_id: str, payload: StatusUpdate, admin: dict = Depends(get_current_admin)):
    if payload.production_status not in PRODUCTION_STATUSES:
        raise HTTPException(400, "Invalid status")
    update = {"production_status": payload.production_status, "updated_at": now_iso()}
    if payload.tracking_number:
        update["tracking_number"] = payload.tracking_number
    await db.orders.update_one(
        {"id": order_id},
        {"$set": update,
         "$push": {"status_history": {"status": payload.production_status, "at": now_iso(), "note": payload.note or ""}}},
    )
    order = await db.orders.find_one({"id": order_id}, {"_id": 0})
    if order:
        if payload.production_status == "delivered":
            await db.albums.update_one({"id": order["album_id"]}, {"$set": {"status": "delivered"}})
        await notify_order_status(order, payload.production_status, payload.note or "")
    return {"success": True}


@router.post("/orders/{order_id}/pdf", status_code=202)
async def generate_pdf(order_id: str, admin: dict = Depends(get_current_admin)):
    """(Re)builds the full print package from the frozen album_snapshot in the BACKGROUND and returns at once:
    Downloads/<order>/Album.pdf + Cover/cover.jpg + Cover/back_cover.jpg + Print/page_NNN.jpg + manifest.json.
    Rendering a 20-sheet album takes well over the 60 s reverse-proxy timeout on the production host, so the admin
    console polls GET /orders/{id}/downloads until package.status is "ready" or "failed"."""
    order = await db.orders.find_one({"id": order_id}, {"_id": 0, "id": 1, "album_snapshot": 1, "production_package": 1})
    if not order:
        raise HTTPException(404, "Order not found")
    if not order.get("album_snapshot"):
        raise HTTPException(400, "This order has no frozen album snapshot to render")
    pkg = order.get("production_package") or {}
    if pkg.get("status") == "building" and not production_build_is_stale(pkg):
        return {"status": "building", "package": pkg, "message": "Production files are already being rendered"}
    schedule_production_package(order_id)
    started = {"status": "building", "started_at": now_iso()}
    return {"status": "building", "package": started, "message": "Rendering started"}


@router.get("/orders/{order_id}/downloads")
async def order_downloads(order_id: str, admin: dict = Depends(get_current_admin)):
    order = await db.orders.find_one({"id": order_id}, {"_id": 0, "order_no": 1, "production_package": 1})
    if not order:
        raise HTTPException(404, "Order not found")
    pkg = order.get("production_package") or {"status": "none"}
    if pkg.get("status") == "building" and production_build_is_stale(pkg):
        # The render never finished (worker restarted / out of memory on the host) — report it instead of spinning forever.
        pkg = {**pkg, "status": "failed", "error": STALE_BUILD_MESSAGE}
        await db.orders.update_one({"id": order_id}, {"$set": {"production_package": pkg}})
    files = list_dir_files(pkg["dir"]) if pkg.get("dir") else []
    return {"package": pkg, "files": [{**f, "url": public_url(f["path"])} for f in files]}


@router.get("/orders/{order_id}/downloads.zip", include_in_schema=False)
async def order_downloads_zip(order_id: str, token: str = ""):
    """Zip of the whole Downloads/<order>/ folder. Token passed as query so it can open in a browser tab."""
    if not token:
        raise HTTPException(401, "Invalid admin token")
    await admin_from_token(token)
    order = await db.orders.find_one({"id": order_id}, {"_id": 0, "order_no": 1, "production_package": 1})
    if not order or not (order.get("production_package") or {}).get("dir"):
        raise HTTPException(404, "Production package not built yet")
    data = await run_in_threadpool(zip_dir, order["production_package"]["dir"])
    if data is None:
        raise HTTPException(404, "Package folder missing")
    name = order["production_package"]["dir"].split("/")[-1]
    return Response(content=data, media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{name}.zip"'})


# ---- Admin-managed images (covers, backgrounds, process bots, promos) ----

@router.post("/images")
async def admin_upload_image(file: UploadFile = File(...), admin: dict = Depends(get_current_admin)):
    """Upload a replacement image for any admin-managed asset. Returns stable URLs (preview + original)."""
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 20MB)")
    try:
        meta = await run_in_threadpool(save_original, "admin", "assets", file.filename or "image.jpg", data)
    except ValueError as e:
        raise HTTPException(400, str(e))
    doc = {"id": new_id(), "filename": file.filename, "uploaded_by": admin["username"], "created_at": now_iso(), **meta,
           "url": public_url(meta["preview_path"]), "thumbnail_url": public_url(meta["thumbnail_path"]),
           "original_url": public_url(meta["original_path"])}
    await db.admin_images.insert_one(dict(doc))
    doc.pop("_id", None)
    return {"image": doc}


VIDEO_TYPES = {"video/mp4": ".mp4", "video/webm": ".webm", "video/quicktime": ".mov"}


@router.post("/videos")
async def admin_upload_video(file: UploadFile = File(...), admin: dict = Depends(get_current_admin)):
    """Upload a short product video for the home sliders (MP4/WebM/MOV, max 60MB). Served from /api/files with
    Range support so it streams on iPhone."""
    ext = VIDEO_TYPES.get((file.content_type or "").lower()) or Path(file.filename or "").suffix.lower()
    if ext not in VIDEO_TYPES.values():
        raise HTTPException(400, "Unsupported video type — upload MP4, WebM or MOV")
    data = await file.read()
    if len(data) > 60 * 1024 * 1024:
        raise HTTPException(400, "Video too large (max 60MB)")
    vid = new_id()
    rel = f"home/videos/{vid}{ext}"
    await run_in_threadpool(_put, rel, data, file.content_type or "video/mp4")
    doc = {"id": vid, "filename": file.filename, "size_bytes": len(data), "uploaded_by": admin["username"],
           "created_at": now_iso(), "path": rel, "url": public_url(rel)}
    await db.admin_videos.insert_one(dict(doc))
    doc.pop("_id", None)
    return {"video": doc}


# ---- Orders export (Excel) ----

EXPORT_RANGES = {"daily": 1, "weekly": 7, "monthly": 31}


@router.get("/orders/export.xlsx")
async def export_orders_xlsx(range: str = "monthly", token: Optional[str] = None, authorization: Optional[str] = Header(None)):
    """Complete order list as an Excel workbook (daily / weekly / monthly / all). Opened via a direct link, so the
    admin JWT may be passed as ?token= like the ZIP download."""
    tok = token or (authorization.split(" ", 1)[1] if authorization and authorization.lower().startswith("bearer ") else None)
    if not tok:
        raise HTTPException(401, "Missing admin token")
    await admin_from_token(tok)
    days = EXPORT_RANGES.get(range)
    q: dict = {}
    if days:
        since = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()
        q = {"created_at": {"$gte": since}}
    orders = await db.orders.find(q, {"_id": 0}).sort("created_at", -1).to_list(5000)
    cust_ids = list({o.get("customer_id") for o in orders if o.get("customer_id")})
    customers = {c["id"]: c for c in await db.customers.find({"id": {"$in": cust_ids}}, {"_id": 0}).to_list(5000)}

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Orders ({range})"
    headers = ["Order ID", "Order date", "Payment date", "Client name", "Mobile", "Email", "GST no.",
               "Album", "Sheets", "Pages", "Photos", "Price per sheet", "Subtotal", "Discount", "Coupon", "Gift wrap",
               "GST", "Total (₹)", "Payment status", "Payment method", "Payment ref", "Production status",
               "Address name", "Address line", "City", "State", "Pincode", "Address phone",
               "Album PDF", "Cover", "Back cover", "Print folder"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B2B2E")
        c.alignment = Alignment(vertical="center")
    for o in orders:
        a = o.get("address") or {}
        cs = o.get("customer_snapshot") or {}
        cust = customers.get(o.get("customer_id"), {})
        snap = o.get("album_snapshot") or {}
        price = o.get("price") or {}
        pkg = o.get("production_package") or {}
        ws.append([
            o.get("order_no"), (o.get("created_at") or "")[:19].replace("T", " "), (o.get("paid_at") or "")[:19].replace("T", " "),
            o.get("client_name") or cs.get("name") or cust.get("name"), cs.get("mobile") or cust.get("mobile"),
            cs.get("email") or cust.get("email"), cust.get("gst_no"),
            snap.get("name") or o.get("album_name"), o.get("sheets"), len(snap.get("pages") or []), len(snap.get("photos") or []),
            price.get("price_per_sheet"), price.get("subtotal"), price.get("discount"), o.get("coupon_code"),
            "Yes" if o.get("gift_wrap") else "No", price.get("gst"), price.get("total"),
            "Payment complete" if o.get("payment_status") == "paid" else "Payment incomplete",
            o.get("payment_method"), o.get("payment_id"), (o.get("production_status") or "").replace("_", " "),
            a.get("name"), ", ".join(filter(None, [a.get("line1"), a.get("line2")])), a.get("city"), a.get("state"),
            a.get("pincode") or a.get("pin"), a.get("phone"),
            pkg.get("pdf_url"), pkg.get("cover_url"), pkg.get("back_cover_url"), (pkg.get("print_urls") or [None])[0],
        ])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 6))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"clickbook-orders-{range}-{datetime.now(timezone.utc).date().isoformat()}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---- Covers ----

@router.post("/covers")
async def create_cover(payload: CoverIn, admin: dict = Depends(get_current_admin)):
    doc = {"id": new_id(), **payload.dict(), "created_at": now_iso()}
    await db.covers.insert_one(dict(doc))
    return {"cover": doc}


@router.put("/covers/{cover_id}")
async def update_cover(cover_id: str, payload: CoverIn, admin: dict = Depends(get_current_admin)):
    await db.covers.update_one({"id": cover_id}, {"$set": payload.dict()})
    return {"success": True}


@router.delete("/covers/{cover_id}")
async def delete_cover(cover_id: str, admin: dict = Depends(get_current_admin)):
    await db.covers.delete_one({"id": cover_id})
    return {"success": True}


# ---- Offers ----

@router.get("/offers")
async def admin_offers(admin: dict = Depends(get_current_admin)):
    return {"offers": await db.offers.find({}, {"_id": 0}).to_list(200)}


@router.post("/offers")
async def create_offer(payload: OfferIn, admin: dict = Depends(get_current_admin)):
    doc = {"id": new_id(), **payload.dict(), "code": payload.code.upper(), "usage_count": 0, "created_at": now_iso()}
    await db.offers.insert_one(dict(doc))
    return {"offer": doc}


@router.put("/offers/{offer_id}")
async def update_offer(offer_id: str, payload: OfferIn, admin: dict = Depends(get_current_admin)):
    update = payload.dict()
    update["code"] = update["code"].upper()
    await db.offers.update_one({"id": offer_id}, {"$set": update})
    return {"success": True}


# ---- Settings / bots / customers ----

@router.put("/settings")
async def update_settings(payload: SettingsUpdate, admin: dict = Depends(get_current_admin)):
    update = {k: v for k, v in payload.dict().items() if v is not None}
    if update:
        await db.settings.update_one({"id": "default"}, {"$set": update}, upsert=True)
    return {"settings": await get_settings_doc()}


@router.get("/process-bots")
async def list_process_bots(admin: dict = Depends(get_current_admin)):
    return {"process_bots": await db.process_bots.find({}, {"_id": 0}).sort("order", 1).to_list(20)}


@router.put("/process-bots/{bot_id}")
async def update_bot(bot_id: str, payload: BotUpdate, admin: dict = Depends(get_current_admin)):
    update = {k: v for k, v in payload.dict().items() if v is not None}
    if update:
        await db.process_bots.update_one({"id": bot_id}, {"$set": update})
    return {"success": True}


@router.get("/customers")
async def admin_customers(admin: dict = Depends(get_current_admin)):
    customers = await db.customers.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)
    pipeline = [{"$group": {"_id": "$customer_id", "orders": {"$sum": 1},
                            "spent": {"$sum": {"$cond": [{"$eq": ["$payment_status", "paid"]}, "$price.total", 0]}}}}]
    stats = {r["_id"]: r async for r in db.orders.aggregate(pipeline)}
    albums = {r["_id"]: r["n"] async for r in db.albums.aggregate([{"$group": {"_id": "$customer_id", "n": {"$sum": 1}}}])}
    for c in customers:
        st = stats.get(c["id"], {})
        c["order_count"] = st.get("orders", 0)
        c["total_spent"] = round(st.get("spent", 0), 2)
        c["album_count"] = albums.get(c["id"], 0)
    return {"customers": customers}


# ---- Layouts & backgrounds (catalog CRUD; GET lists live in routers/catalog.py) ----

class LayoutIn(BaseModel):
    name: str
    photo_count: int
    positions: List[Dict[str, Any]]
    active: bool = True


class BackgroundIn(BaseModel):
    name: str
    color: str
    image_url: Optional[str] = None
    active: bool = True


@router.get("/layouts")
async def admin_layouts(admin: dict = Depends(get_current_admin)):
    return {"layouts": await db.layouts.find({}, {"_id": 0}).sort("photo_count", 1).to_list(100)}


@router.post("/layouts")
async def create_layout(payload: LayoutIn, admin: dict = Depends(get_current_admin)):
    if payload.photo_count != len(payload.positions):
        raise HTTPException(400, "positions must have exactly photo_count rectangles")
    doc = {"id": new_id(), **payload.dict(), "created_at": now_iso()}
    await db.layouts.insert_one(dict(doc))
    return {"layout": doc}


@router.put("/layouts/{layout_id}")
async def update_layout(layout_id: str, payload: LayoutIn, admin: dict = Depends(get_current_admin)):
    await db.layouts.update_one({"id": layout_id}, {"$set": payload.dict()})
    return {"success": True}


@router.delete("/layouts/{layout_id}")
async def delete_layout(layout_id: str, admin: dict = Depends(get_current_admin)):
    await db.layouts.delete_one({"id": layout_id})
    return {"success": True}


@router.get("/backgrounds")
async def admin_backgrounds(admin: dict = Depends(get_current_admin)):
    return {"backgrounds": await db.backgrounds.find({}, {"_id": 0}).to_list(200)}


@router.post("/backgrounds")
async def create_background(payload: BackgroundIn, admin: dict = Depends(get_current_admin)):
    doc = {"id": new_id(), **payload.dict(), "created_at": now_iso()}
    await db.backgrounds.insert_one(dict(doc))
    return {"background": doc}


@router.put("/backgrounds/{bg_id}")
async def update_background(bg_id: str, payload: BackgroundIn, admin: dict = Depends(get_current_admin)):
    await db.backgrounds.update_one({"id": bg_id}, {"$set": payload.dict()})
    return {"success": True}


@router.delete("/backgrounds/{bg_id}")
async def delete_background(bg_id: str, admin: dict = Depends(get_current_admin)):
    await db.backgrounds.delete_one({"id": bg_id})
    return {"success": True}
